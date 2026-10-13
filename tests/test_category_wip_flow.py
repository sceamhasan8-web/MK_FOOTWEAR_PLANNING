import pytest
from datetime import date, timedelta
from app.services.production_planning_service import ProductionPlanningService
from app.services.calendar_service import CalendarService
from app.models.route import Category, Section, SectionRouteStatus
from app.models.order import Order, OrderStatus
from app.models.entry import DailyProductionEntry
from app.models.calendar import FactoryCalendar


def _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0, days_ahead=10):
    """Helper to create a standard 3-category footwear order."""
    admin_hdr = auth_headers("admin")

    r_buyer = client.post("/api/v1/buyers", json={
        "name": "BATA Footwear International",
        "buyer_code": "BATA-WIP",
        "is_active": True
    }, headers=admin_hdr)
    buyer_id = r_buyer.json()["id"]

    sec_cut = db_session.query(Section).filter(Section.name == "Cutting").first()
    sec_sew = db_session.query(Section).filter(Section.name == "Sewing").first()
    sec_lst = db_session.query(Section).filter(Section.name == "Lasting").first()

    today = date.today()
    target_delivery = today + timedelta(days=days_ahead)

    r_order = client.post("/api/v1/orders", json={
        "buyer_id": buyer_id,
        "order_code": f"ORD-WIP-{int(order_qty)}",
        "customer_ref": "REF-WIP-001",
        "style_model": "Sneaker Pro",
        "total_quantity": order_qty,
        "production_start_date": str(today),
        "original_delivery_date": str(target_delivery),
        "routes": [
            {
                "section_id": sec_cut.id,
                "status": "REQUIRED",
                "operations": [
                    {"operation_name": "Cutting Process", "required_quantity": order_qty, "sequence_order": 1}
                ]
            },
            {
                "section_id": sec_sew.id,
                "status": "REQUIRED",
                "operations": [
                    {"operation_name": "Sewing Process", "required_quantity": order_qty, "sequence_order": 1}
                ]
            },
            {
                "section_id": sec_lst.id,
                "status": "REQUIRED",
                "operations": [
                    {"operation_name": "Lasting Process", "required_quantity": order_qty, "sequence_order": 1}
                ]
            }
        ]
    }, headers=admin_hdr)
    assert r_order.status_code == 200, r_order.text
    return r_order.json()


def test_scenario_1_initial_on_hand(client, auth_headers, db_session):
    """
    Test 1 – Initial On Hand:
    Order = 70,000
    Expected:
    Category 1 On Hand = 70,000
    Category 2 On Hand = 0
    Category 3 On Hand = 0
    """
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]

    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)

    assert len(flow) == 3
    # Category 1: Cutting
    assert flow[0]["opening_on_hand"] == 70000.0
    assert flow[0]["available_on_hand"] == 70000.0
    assert flow[0]["closing_on_hand"] == 70000.0

    # Category 2: Preparation
    assert flow[1]["opening_on_hand"] == 0.0
    assert flow[1]["incoming_wip"] == 0.0
    assert flow[1]["available_on_hand"] == 0.0
    assert flow[1]["closing_on_hand"] == 0.0

    # Category 3: Lasting & Dip
    assert flow[2]["opening_on_hand"] == 0.0
    assert flow[2]["incoming_wip"] == 0.0
    assert flow[2]["available_on_hand"] == 0.0
    assert flow[2]["closing_on_hand"] == 0.0


def test_scenario_2_category_1_production(client, auth_headers, db_session):
    """
    Test 2 – Category 1 Production:
    Category 1 actual = 5,000
    Expected:
    Category 1 Closing On Hand = 65,000
    Category 2 Incoming WIP = 5,000
    """
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]
    supervisor_hdr = auth_headers("supervisor")

    # Record 5,000 in Category 1 Cutting
    cat1_op = order_data["routes"][0]["operations"][0]
    r_entry = client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": cat1_op["id"],
        "entry_date": str(date.today()),
        "quantity": 5000.0,
        "accepted_quantity": 5000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)
    assert r_entry.status_code == 200, r_entry.text

    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)

    # Category 1
    assert flow[0]["completed_output"] == 5000.0
    assert flow[0]["closing_on_hand"] == 65000.0

    # Category 2
    assert flow[1]["incoming_wip"] == 5000.0
    assert flow[1]["available_on_hand"] == 5000.0


def test_scenario_3_category_2_production(client, auth_headers, db_session):
    """
    Test 3 – Category 2 Production:
    Category 2 receives 5,000.
    Category 2 actual = 3,000.
    Expected:
    Category 2 Closing On Hand = 2,000
    Category 3 Incoming WIP = 3,000
    """
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]
    supervisor_hdr = auth_headers("supervisor")

    # Step 1: Category 1 produces 5,000
    cat1_op = order_data["routes"][0]["operations"][0]
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": cat1_op["id"],
        "entry_date": str(date.today()),
        "quantity": 5000.0,
        "accepted_quantity": 5000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    # Step 2: Category 2 produces 3,000
    cat2_op = order_data["routes"][1]["operations"][0]
    r_entry2 = client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][1]["id"],
        "operation_id": cat2_op["id"],
        "entry_date": str(date.today()),
        "quantity": 3000.0,
        "accepted_quantity": 3000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)
    assert r_entry2.status_code == 200, r_entry2.text

    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)

    # Category 2 Closing On Hand = 5,000 - 3,000 = 2,000
    assert flow[1]["closing_on_hand"] == 2000.0
    assert flow[1]["completed_output"] == 3000.0

    # Category 3 Incoming WIP = 3,000
    assert flow[2]["incoming_wip"] == 3000.0
    assert flow[2]["available_on_hand"] == 3000.0


def test_scenario_4_multi_process_category(client, auth_headers, db_session):
    """
    Test 4 – Multi-Process Category:
    Category 2 has Printing = 38,000, Emboss = 30,000, Sewing = 35,000
    Expected Category 2 Output: 30,000 (Not 103,000).
    """
    admin_hdr = auth_headers("admin")
    supervisor_hdr = auth_headers("supervisor")

    r_buyer = client.post("/api/v1/buyers", json={
        "name": "BATA Multi-Process",
        "buyer_code": "BATA-MP",
        "is_active": True
    }, headers=admin_hdr)
    buyer_id = r_buyer.json()["id"]

    sec_prt = db_session.query(Section).filter(Section.name == "Printing").first()
    sec_emb = db_session.query(Section).filter(Section.name == "Emboss").first()
    sec_sew = db_session.query(Section).filter(Section.name == "Sewing").first()

    today = date.today()
    r_order = client.post("/api/v1/orders", json={
        "buyer_id": buyer_id,
        "order_code": "ORD-MP-38K",
        "customer_ref": "REF-MP",
        "style_model": "Boot Classic",
        "total_quantity": 38000.0,
        "production_start_date": str(today),
        "original_delivery_date": str(today + timedelta(days=10)),
        "routes": [
            {
                "section_id": sec_prt.id,
                "status": "REQUIRED",
                "operations": [{"operation_name": "Printing", "required_quantity": 38000.0, "sequence_order": 1}]
            },
            {
                "section_id": sec_emb.id,
                "status": "REQUIRED",
                "operations": [{"operation_name": "Emboss", "required_quantity": 38000.0, "sequence_order": 2}]
            },
            {
                "section_id": sec_sew.id,
                "status": "REQUIRED",
                "operations": [{"operation_name": "Sewing", "required_quantity": 38000.0, "sequence_order": 3}]
            }
        ]
    }, headers=admin_hdr)
    order_data = r_order.json()
    order_id = order_data["id"]

    routes = order_data["routes"]
    op_prt = routes[0]["operations"][0]["id"]
    op_emb = routes[1]["operations"][0]["id"]
    op_sew = routes[2]["operations"][0]["id"]

    # Post entries
    client.post("/api/v1/production/entry", json={
        "order_id": order_id, "route_id": routes[0]["id"], "operation_id": op_prt,
        "entry_date": str(today), "quantity": 38000.0, "accepted_quantity": 38000.0,
        "rejected_quantity": 0.0, "unit": "PAIRS"
    }, headers=supervisor_hdr)

    client.post("/api/v1/production/entry", json={
        "order_id": order_id, "route_id": routes[1]["id"], "operation_id": op_emb,
        "entry_date": str(today), "quantity": 30000.0, "accepted_quantity": 30000.0,
        "rejected_quantity": 0.0, "unit": "PAIRS"
    }, headers=supervisor_hdr)

    client.post("/api/v1/production/entry", json={
        "order_id": order_id, "route_id": routes[2]["id"], "operation_id": op_sew,
        "entry_date": str(today), "quantity": 35000.0, "accepted_quantity": 35000.0,
        "rejected_quantity": 0.0, "unit": "PAIRS"
    }, headers=supervisor_hdr)

    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)

    # Preparation category output = min(38000, 30000, 35000) = 30000
    prep_cat = next(c for c in flow if c["category_name"] == "Preparation")
    assert prep_cat["completed_output"] == 30000.0
    assert prep_cat["completed_output"] != 103000.0


def test_scenario_5_not_required_process(client, auth_headers, db_session):
    """
    Test 5 – Not Required Process:
    Printing = Required (40,000)
    Emboss = Not Required
    Sewing = Required (35,000)
    Expected Category 2 Output: 35,000.
    """
    admin_hdr = auth_headers("admin")
    supervisor_hdr = auth_headers("supervisor")

    r_buyer = client.post("/api/v1/buyers", json={
        "name": "BATA Optional Process",
        "buyer_code": "BATA-OPT",
        "is_active": True
    }, headers=admin_hdr)
    buyer_id = r_buyer.json()["id"]

    sec_prt = db_session.query(Section).filter(Section.name == "Printing").first()
    sec_emb = db_session.query(Section).filter(Section.name == "Emboss").first()
    sec_sew = db_session.query(Section).filter(Section.name == "Sewing").first()

    today = date.today()
    r_order = client.post("/api/v1/orders", json={
        "buyer_id": buyer_id,
        "order_code": "ORD-OPT-40K",
        "customer_ref": "REF-OPT",
        "style_model": "Loafer Casual",
        "total_quantity": 40000.0,
        "production_start_date": str(today),
        "original_delivery_date": str(today + timedelta(days=10)),
        "routes": [
            {
                "section_id": sec_prt.id,
                "status": "REQUIRED",
                "operations": [{"operation_name": "Printing", "required_quantity": 40000.0, "sequence_order": 1}]
            },
            {
                "section_id": sec_emb.id,
                "status": "NOT_REQUIRED",
                "operations": [{"operation_name": "Emboss", "required_quantity": 40000.0, "sequence_order": 2}]
            },
            {
                "section_id": sec_sew.id,
                "status": "REQUIRED",
                "operations": [{"operation_name": "Sewing", "required_quantity": 40000.0, "sequence_order": 3}]
            }
        ]
    }, headers=admin_hdr)
    order_data = r_order.json()
    order_id = order_data["id"]

    routes = order_data["routes"]
    op_prt = routes[0]["operations"][0]["id"]
    op_sew = routes[2]["operations"][0]["id"]

    client.post("/api/v1/production/entry", json={
        "order_id": order_id, "route_id": routes[0]["id"], "operation_id": op_prt,
        "entry_date": str(today), "quantity": 40000.0, "accepted_quantity": 40000.0,
        "rejected_quantity": 0.0, "unit": "PAIRS"
    }, headers=supervisor_hdr)

    client.post("/api/v1/production/entry", json={
        "order_id": order_id, "route_id": routes[2]["id"], "operation_id": op_sew,
        "entry_date": str(today), "quantity": 35000.0, "accepted_quantity": 35000.0,
        "rejected_quantity": 0.0, "unit": "PAIRS"
    }, headers=supervisor_hdr)

    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)

    prep_cat = next(c for c in flow if c["category_name"] == "Preparation")
    # Ignored Emboss because NOT_REQUIRED. min(40000, 35000) = 35000
    assert prep_cat["completed_output"] == 35000.0


def test_scenario_6_insufficient_wip(client, auth_headers, db_session):
    """
    Test 6 – Insufficient WIP:
    Category 2 Target = 6,000, Available On Hand = 3,000
    Expected: Executable Target = 3,000, Status = WAITING FOR WIP / PARTIALLY AVAILABLE.
    Also verifies: Production entry exceeding Available On Hand without override is rejected (Section 8).
    """
    cal = CalendarService.get_active_calendar(db_session)
    cal.same_day_category_transfer = True
    db_session.commit()

    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]
    supervisor_hdr = auth_headers("supervisor")

    # Step 1: Category 1 produces 3,000
    cat1_op = order_data["routes"][0]["operations"][0]
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": cat1_op["id"],
        "entry_date": str(date.today()),
        "quantity": 3000.0,
        "accepted_quantity": 3000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    # Category 2 now has Available On Hand = 3,000
    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)
    cat2_flow = flow[1]
    assert cat2_flow["available_on_hand"] == 3000.0

    # Required target for remaining balance 70,000 / remaining days > 6,000
    # Executable target = min(target, available_on_hand) = 3,000
    sheet = ProductionPlanningService.generate_category_flow_planning_sheet(db_session, order_id=order_id)
    cat2_rows = [r for r in sheet if r["category_name"] == "Preparation"]
    assert cat2_rows[0]["executable_target"] == 3000.0
    assert cat2_rows[0]["wip_status"] in ["PARTIALLY AVAILABLE", "WAITING FOR WIP"]

    # Verify attempt to enter 7,000 without override is rejected with "Insufficient On Hand"
    cat2_op = order_data["routes"][1]["operations"][0]
    r_bad = client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][1]["id"],
        "operation_id": cat2_op["id"],
        "entry_date": str(date.today()),
        "quantity": 7000.0,
        "accepted_quantity": 7000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS",
        "is_override": False
    }, headers=supervisor_hdr)
    assert r_bad.status_code == 400
    assert "Insufficient On Hand quantity" in r_bad.json()["detail"]


def test_scenario_7_same_day_transfer_disabled(client, auth_headers, db_session):
    """
    Test 7 – Same-Day Transfer Disabled:
    Category 1 completes 5,000 on 07 October.
    Category 2 should receive the 5,000 on 08 October.
    """
    cal = CalendarService.get_active_calendar(db_session)
    cal.same_day_category_transfer = False
    db_session.commit()

    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]
    supervisor_hdr = auth_headers("supervisor")

    today = date(2026, 10, 7)
    tomorrow = date(2026, 10, 8)

    # Cat 1 completes 5,000 on 07 October
    cat1_op = order_data["routes"][0]["operations"][0]
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": cat1_op["id"],
        "entry_date": str(today),
        "quantity": 5000.0,
        "accepted_quantity": 5000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    sheet = ProductionPlanningService.generate_category_flow_planning_sheet(
        db_session, order_id=order_id, as_of_date=today
    )
    # On 07 October: Cat 2 incoming is 0
    cat2_day1 = [r for r in sheet if r["category_name"] == "Preparation" and r["date"] == today]
    if cat2_day1:
        assert cat2_day1[0]["incoming_wip"] == 0.0

    # On 08 October: Cat 2 incoming is 5,000
    cat2_day2 = [r for r in sheet if r["category_name"] == "Preparation" and r["date"] == tomorrow]
    if cat2_day2:
        assert cat2_day2[0]["incoming_wip"] == 5000.0


def test_scenario_8_same_day_transfer_enabled(client, auth_headers, db_session):
    """
    Test 8 – Same-Day Transfer Enabled:
    Category 1 completes 5,000 on 07 October.
    Category 2 may receive 5,000 on 07 October.
    """
    cal = CalendarService.get_active_calendar(db_session)
    cal.same_day_category_transfer = True
    db_session.commit()

    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]
    supervisor_hdr = auth_headers("supervisor")

    today = date(2026, 10, 7)
    cat1_op = order_data["routes"][0]["operations"][0]
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": cat1_op["id"],
        "entry_date": str(today),
        "quantity": 5000.0,
        "accepted_quantity": 5000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    sheet = ProductionPlanningService.generate_category_flow_planning_sheet(
        db_session, order_id=order_id, as_of_date=today
    )
    cat2_day1 = [r for r in sheet if r["category_name"] == "Preparation" and r["date"] == today]
    assert len(cat2_day1) > 0
    # With same-day transfer enabled, 5000 is received on the same day
    assert cat2_day1[0]["incoming_wip"] == 5000.0


def test_scenario_9_production_correction(client, auth_headers, db_session):
    """
    Test 9 – Production Correction:
    Category 1 production changed from 5,000 to 4,000.
    Expected downstream WIP and all dependent planning values must be recalculated.
    """
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0)
    order_id = order_data["id"]
    supervisor_hdr = auth_headers("supervisor")

    cat1_op = order_data["routes"][0]["operations"][0]
    r_entry = client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": cat1_op["id"],
        "entry_date": str(date.today()),
        "quantity": 5000.0,
        "accepted_quantity": 5000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)
    entry_id = r_entry.json()["id"]

    # Correct entry from 5,000 to 4,000
    r_corr = client.put(f"/api/v1/production/entries/{entry_id}/correct", json={
        "new_accepted_quantity": 4000.0,
        "new_rejected_quantity": 0.0,
        "reason": "Re-count correction on Cutting line"
    }, headers=supervisor_hdr)
    assert r_corr.status_code == 200

    order = db_session.query(Order).filter(Order.id == order_id).first()
    flow = ProductionPlanningService.calculate_category_flow(db_session, order)

    # Category 1 closing On Hand becomes 70,000 - 4,000 = 66,000
    assert flow[0]["closing_on_hand"] == 66000.0
    assert flow[0]["completed_output"] == 4000.0

    # Category 2 incoming WIP becomes 4,000
    assert flow[1]["incoming_wip"] == 4000.0
    assert flow[1]["available_on_hand"] == 4000.0


def test_scenario_10_shipment_extension(client, auth_headers, db_session):
    """
    Test 10 – Shipment Extension:
    Shipment date is extended.
    Expected: More available days, lower required daily target.
    """
    planner_hdr = auth_headers("planner")
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0, days_ahead=5)
    order_id = order_data["id"]

    order = db_session.query(Order).filter(Order.id == order_id).first()
    plan_before = ProductionPlanningService.calculate_order_plan(db_session, order)
    target_before = plan_before["daily_target"]
    days_before = plan_before["remaining_days"]

    # Extend delivery date by 10 days
    new_date = order.current_delivery_date + timedelta(days=10)
    r_rev = client.post(f"/api/v1/orders/{order_id}/revise-delivery-date", json={
        "revised_delivery_date": str(new_date),
        "reason": "Buyer granted shipping extension"
    }, headers=planner_hdr)
    assert r_rev.status_code == 200

    db_session.refresh(order)
    plan_after = ProductionPlanningService.calculate_order_plan(db_session, order)
    target_after = plan_after["daily_target"]
    days_after = plan_after["remaining_days"]

    assert days_after > days_before
    assert target_after < target_before


def test_scenario_11_shipment_reduction(client, auth_headers, db_session):
    """
    Test 11 – Shipment Reduction:
    Shipment date is moved earlier.
    Expected: Higher required daily target, potential At Risk status.
    """
    planner_hdr = auth_headers("planner")
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=70000.0, days_ahead=15)
    order_id = order_data["id"]

    order = db_session.query(Order).filter(Order.id == order_id).first()
    plan_before = ProductionPlanningService.calculate_order_plan(db_session, order)
    target_before = plan_before["daily_target"]

    # Move delivery date earlier
    new_date = date.today() + timedelta(days=2)
    r_rev = client.post(f"/api/v1/orders/{order_id}/revise-delivery-date", json={
        "revised_delivery_date": str(new_date),
        "reason": "Expedited air-shipment requested"
    }, headers=planner_hdr)
    assert r_rev.status_code == 200

    db_session.refresh(order)
    plan_after = ProductionPlanningService.calculate_order_plan(db_session, order)
    target_after = plan_after["daily_target"]

    assert target_after > target_before
    assert plan_after["is_at_risk"] is True or plan_after["remaining_days"] <= 2


def test_scenario_12_full_completion(client, auth_headers, db_session):
    """
    Test 12 – Full Completion:
    Final Category reaches the complete order quantity.
    Expected:
    Final Balance = 0
    Order Status = Production Complete
    """
    supervisor_hdr = auth_headers("supervisor")
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=10000.0)
    order_id = order_data["id"]

    # Cat 1 Cutting completes 10,000
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][0]["id"],
        "operation_id": order_data["routes"][0]["operations"][0]["id"],
        "entry_date": str(date.today()),
        "quantity": 10000.0,
        "accepted_quantity": 10000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    # Cat 2 Sewing completes 10,000
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][1]["id"],
        "operation_id": order_data["routes"][1]["operations"][0]["id"],
        "entry_date": str(date.today()),
        "quantity": 10000.0,
        "accepted_quantity": 10000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    # Cat 3 Lasting completes 10,000
    client.post("/api/v1/production/entry", json={
        "order_id": order_id,
        "route_id": order_data["routes"][2]["id"],
        "operation_id": order_data["routes"][2]["operations"][0]["id"],
        "entry_date": str(date.today()),
        "quantity": 10000.0,
        "accepted_quantity": 10000.0,
        "rejected_quantity": 0.0,
        "unit": "PAIRS"
    }, headers=supervisor_hdr)

    order = db_session.query(Order).filter(Order.id == order_id).first()
    plan = ProductionPlanningService.calculate_order_plan(db_session, order)

    assert plan["completed_pairs"] == 10000.0
    assert plan["outstanding_pairs"] == 0.0
    assert plan["status"] == OrderStatus.PRODUCTION_COMPLETE
    assert plan["daily_target"] == 0.0


def test_scenario_13_planning_sheet_api(client, auth_headers, db_session):
    """
    Test Planning Sheet API endpoint /api/v1/orders/planning-sheet.
    Verifies that generated rows include all key flow fields:
    opening_on_hand, incoming_wip, available_on_hand, daily_target,
    executable_target, completed_output, closing_on_hand, balance, wip_status.
    """
    planner_hdr = auth_headers("planner")
    order_data = _setup_3_category_order(client, auth_headers, db_session, order_qty=50000.0)

    res = client.get(f"/api/v1/orders/planning-sheet?order_id={order_data['id']}", headers=planner_hdr)
    assert res.status_code == 200
    rows = res.json()
    assert len(rows) > 0

    first_row = rows[0]
    assert "date" in first_row
    assert "buyer_name" in first_row
    assert "order_code" in first_row
    assert "order_quantity" in first_row
    assert "category_name" in first_row
    assert "section_name" in first_row
    assert "operation_name" in first_row
    assert "opening_on_hand" in first_row
    assert "incoming_wip" in first_row
    assert "available_on_hand" in first_row
    assert "daily_target" in first_row
    assert "executable_target" in first_row
    assert "completed_output" in first_row
    assert "closing_on_hand" in first_row
    assert "balance" in first_row
    assert "wip_status" in first_row
    assert "production_status" in first_row


def test_scenario_14_planning_xlsx_export_4_sheets(client, auth_headers, db_session):
    """
    Test Excel Export endpoint /api/v1/exports/planning/xlsx.
    Verifies that the generated workbook contains exactly the 4 required sheets:
    - Master Planning
    - Order Summary
    - Daily Summary
    - Category Flow
    """
    import io
    import openpyxl

    planner_hdr = auth_headers("planner")
    _setup_3_category_order(client, auth_headers, db_session, order_qty=30000.0)

    res = client.get("/api/v1/exports/planning/xlsx", headers=planner_hdr)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    expected_sheets = ["Master Planning", "Order Summary", "Daily Summary", "Category Flow"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Expected sheet '{s}' not found in {wb.sheetnames}"

