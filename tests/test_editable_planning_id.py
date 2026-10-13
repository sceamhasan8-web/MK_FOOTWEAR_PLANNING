import io
import csv
import openpyxl
import pytest
from datetime import date, timedelta
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.database import get_db
from app.models.user import User, UserRole
from app.models.buyer import Buyer
from app.models.order import Order, OrderStatus
from app.models.route import Category, Section, OrderRoute, OrderOperation, SectionRouteStatus, ProductionTargetSetting
from app.models.entry import DailyProductionEntry
from app.services.calendar_service import CalendarService
from app.services.production_planning_service import ProductionPlanningService
from app.services.export_service import ExportService
from app.services.settings_service import SettingsService
from app.api.deps import create_access_token


@pytest.fixture
def client_and_db():
    client = TestClient(app)
    db = next(get_db())
    yield client, db


@pytest.fixture
def admin_headers(client_and_db):
    _, db = client_and_db
    admin = db.query(User).filter(User.role == UserRole.ADMIN).first()
    if not admin:
        admin = User(
            username="admin_planning_id",
            email="admin_pid@factory.com",
            hashed_password="pw",
            role=UserRole.ADMIN,
            is_active=True
        )
        db.add(admin)
        db.commit()
        db.refresh(admin)
    token = create_access_token(data={"sub": admin.username, "user_id": admin.id, "role": admin.role.value})
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def footwear_planning_env(client_and_db):
    _, db = client_and_db
    cal = CalendarService.get_active_calendar(db)

    cat_cut = db.query(Category).filter(Category.name == "Cat Cut Test").first()
    if not cat_cut:
        cat_cut = Category(name="Cat Cut Test", display_order=1, is_active=True)
        db.add(cat_cut)

    cat_prep = db.query(Category).filter(Category.name == "Cat Prep Test").first()
    if not cat_prep:
        cat_prep = Category(name="Cat Prep Test", display_order=2, is_active=True)
        db.add(cat_prep)

    cat_fin = db.query(Category).filter(Category.name == "Cat Fin Test").first()
    if not cat_fin:
        cat_fin = Category(name="Cat Fin Test", display_order=3, is_active=True)
        db.add(cat_fin)
    db.commit()

    buyer = db.query(Buyer).filter(Buyer.buyer_code == "PID-BUYER").first()
    if not buyer:
        buyer = Buyer(name="Planning ID Buyer", buyer_code="PID-BUYER", address="Export City", is_active=True)
        db.add(buyer)
        db.commit()
        db.refresh(buyer)

    return {
        "cal": cal,
        "buyer": buyer,
        "categories": {"cut": cat_cut, "prep": cat_prep, "fin": cat_fin}
    }


def reset_test_sections(db, cats):
    # Clean up any leftover PID sections beyond the standard 4
    extras = db.query(Section).filter(
        Section.name.like("PID%"),
        ~Section.name.in_(["PID Cutting", "PID Printing", "PID Logo Printing", "PID Stitching", "PID Lasting"])
    ).all()
    for ex in extras:
        routes = db.query(OrderRoute).filter(OrderRoute.section_id == ex.id).all()
        for r in routes:
            db.query(DailyProductionEntry).filter(DailyProductionEntry.route_id == r.id).delete(synchronize_session=False)
            db.query(OrderOperation).filter(OrderOperation.route_id == r.id).delete(synchronize_session=False)
            db.delete(r)
        db.query(ProductionTargetSetting).filter(ProductionTargetSetting.section_id == ex.id).delete(synchronize_session=False)
        db.delete(ex)
    db.commit()

    # Find existing
    s_cut = db.query(Section).filter(Section.name == "PID Cutting").first()
    s_print = db.query(Section).filter(Section.name.in_(["PID Printing", "PID Logo Printing"])).first()
    s_stitch = db.query(Section).filter(Section.name == "PID Stitching").first()
    s_last = db.query(Section).filter(Section.name == "PID Lasting").first()

    # Step 1: Temporarily displace planning IDs to avoid intermediate collisions
    for idx, s in enumerate([s_cut, s_print, s_stitch, s_last]):
        if s:
            s.planning_id = 900 + idx
            s.display_order = 900 + idx
    db.commit()

    # Step 2: Assign desired test state
    if not s_cut:
        s_cut = Section(name="PID Cutting", category_id=cats["cut"].id, planning_id=101, display_order=101, is_active=True)
        db.add(s_cut)
    else:
        s_cut.name = "PID Cutting"
        s_cut.planning_id = 101
        s_cut.display_order = 101
        s_cut.is_active = True

    if not s_print:
        s_print = Section(name="PID Printing", category_id=cats["prep"].id, planning_id=104, display_order=104, is_active=True)
        db.add(s_print)
    else:
        s_print.name = "PID Printing"
        s_print.planning_id = 104
        s_print.display_order = 104
        s_print.is_active = True

    if not s_stitch:
        s_stitch = Section(name="PID Stitching", category_id=cats["prep"].id, planning_id=102, display_order=102, is_active=True)
        db.add(s_stitch)
    else:
        s_stitch.name = "PID Stitching"
        s_stitch.planning_id = 102
        s_stitch.display_order = 102
        s_stitch.is_active = True

    if not s_last:
        s_last = Section(name="PID Lasting", category_id=cats["fin"].id, planning_id=103, display_order=103, is_active=True)
        db.add(s_last)
    else:
        s_last.name = "PID Lasting"
        s_last.planning_id = 103
        s_last.display_order = 103
        s_last.is_active = True

    db.commit()
    for s in [s_cut, s_print, s_stitch, s_last]:
        db.refresh(s)
    return {"cut": s_cut, "print": s_print, "stitch": s_stitch, "last": s_last}


def create_test_order(db, buyer, sections, order_code="ORD-PID-TEST", total_qty=20000.0, days_ahead=10):
    existing = db.query(Order).filter(Order.order_code == order_code).first()
    if existing:
        db.query(DailyProductionEntry).filter(DailyProductionEntry.order_id == existing.id).delete(synchronize_session=False)
        route_ids = [r.id for r in existing.routes]
        if route_ids:
            db.query(OrderOperation).filter(OrderOperation.route_id.in_(route_ids)).delete(synchronize_session=False)
            db.query(OrderRoute).filter(OrderRoute.id.in_(route_ids)).delete(synchronize_session=False)
        db.delete(existing)
        db.commit()

    user = db.query(User).first()
    user_id = user.id if user else 1

    start_d = date.today()
    ship_d = date.today() + timedelta(days=days_ahead)

    order = Order(
        order_code=order_code,
        buyer_id=buyer.id,
        customer_ref=f"REF-{order_code}",
        po_number=f"PO-{order_code}",
        style_model="Running PID Edition",
        total_quantity=total_qty,
        original_order_date=start_d,
        production_start_date=start_d,
        original_delivery_date=ship_d,
        current_delivery_date=ship_d,
        status=OrderStatus.IN_PRODUCTION,
        created_by_id=user_id,
    )
    db.add(order)
    db.flush()

    # Create routes in physical workflow sequence (Cutting -> Printing -> Stitching -> Lasting)
    for seq, sec in enumerate(sections, 1):
        rt = OrderRoute(
            order_id=order.id,
            section_id=sec.id,
            sequence_order=seq,
            status=SectionRouteStatus.REQUIRED,
            has_components=False
        )
        db.add(rt)
        db.flush()

        op = OrderOperation(
            route_id=rt.id,
            operation_name=f"{sec.name} Main",
            sequence_order=1,
            unit="PAIRS",
            required_quantity=total_qty
        )
        db.add(op)
    db.commit()
    db.refresh(order)
    return order


def test_a_planning_id_determines_column_order(client_and_db, footwear_planning_env):
    """
    Test A: Planning ID Determines Column Order
    Create active sections:
    - Cutting: Planning ID 101
    - Printing: Planning ID 104
    - Stitching: Planning ID 102
    - Lasting: Planning ID 103
    Expected table order: Cutting -> Stitching -> Lasting -> Printing
    """
    _, db = client_and_db
    cats = footwear_planning_env["categories"]
    buyer = footwear_planning_env["buyer"]

    sec_dict = reset_test_sections(db, cats)
    sec_objs = [sec_dict["cut"], sec_dict["print"], sec_dict["stitch"], sec_dict["last"]]

    order = create_test_order(db, buyer, sec_objs, order_code="ORD-TEST-A-PID")

    sheet = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)
    sec_names = [s["section_name"] for s in sheet["sections"]]
    pids = [s["planning_id"] for s in sheet["sections"]]

    # Verify strictly ascending planning_id: 101, 102, 103, 104
    assert pids == [101, 102, 103, 104], f"Expected Planning IDs [101, 102, 103, 104], got {pids}"
    assert sec_names == ["PID CUTTING", "PID STITCHING", "PID LASTING", "PID PRINTING"], (
        f"Expected Cutting -> Stitching -> Lasting -> Printing, got {sec_names}"
    )

    # Verify summary keys order matches sheet sections order
    summary_keys = list(sheet["summary"].keys())
    assert summary_keys == sec_names


def test_b_change_planning_id(client_and_db, admin_headers, footwear_planning_env):
    """
    Test B: Change Planning ID
    Change Printing from Planning ID 104 to Planning ID 102 after reordering affected sections.
    Expected:
    - Configuration saves.
    - Table reorders after refresh.
    - Printing's actual production remains unchanged.
    - Printing's target settings remain associated with its internal ID.
    """
    client, db = client_and_db
    buyer = footwear_planning_env["buyer"]
    cats = footwear_planning_env["categories"]

    sec_dict = reset_test_sections(db, cats)
    s_cut = sec_dict["cut"]
    s_print = sec_dict["print"]
    s_stitch = sec_dict["stitch"]
    s_last = sec_dict["last"]

    # Configure target settings attached to internal IDs
    SettingsService.update_section(db, s_print.id, s_print.name, s_print.category_id, planning_id=104, min_target=1200, max_target=3000)
    SettingsService.update_section(db, s_stitch.id, s_stitch.name, s_stitch.category_id, planning_id=102, min_target=1800, max_target=4500)

    # Create order with some production recorded for Printing
    order = create_test_order(db, buyer, [s_cut, s_print, s_stitch, s_last], order_code="ORD-TEST-B-PID")
    rt_print = [r for r in order.routes if r.section_id == s_print.id][0]
    entry = DailyProductionEntry(
        order_id=order.id,
        route_id=rt_print.id,
        operation_id=rt_print.operations[0].id,
        entry_date=date.today(),
        quantity=500.0,
        accepted_quantity=480.0,
        rejected_quantity=20.0,
        entered_by_id=1
    )
    db.add(entry)
    db.commit()

    # Controlled batch reorder: Stitching 102 -> 104, Printing 104 -> 102
    reorder_payload = {
        "items": [
            {"section_id": s_print.id, "planning_id": 102},
            {"section_id": s_stitch.id, "planning_id": 104}
        ],
        "reason": "Swap printing and stitching display positions"
    }
    resp = client.put("/api/v1/settings/sections/reorder", json=reorder_payload, headers=admin_headers)
    assert resp.status_code == 200, resp.text
    db.expire_all()

    # Verify table reorders after refresh
    sheet = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)
    sec_names = [s["section_name"] for s in sheet["sections"]]
    pids = [s["planning_id"] for s in sheet["sections"]]

    # Now: Cutting (101) -> Printing (102) -> Lasting (103) -> Stitching (104)
    assert pids == [101, 102, 103, 104]
    assert sec_names == ["PID CUTTING", "PID PRINTING", "PID LASTING", "PID STITCHING"]

    # Verify Printing's target settings remain linked to internal ID (min 1200, max 3000)
    ts_print = db.query(ProductionTargetSetting).filter(ProductionTargetSetting.section_id == s_print.id).first()
    assert ts_print.minimum_daily_target == 1200
    assert ts_print.maximum_daily_target == 3000

    # Verify Printing's recorded production remains 480 accepted pairs
    assert sheet["summary"]["PID PRINTING"]["total_ot"] == 480.0


def test_c_duplicate_planning_id(client_and_db, admin_headers, footwear_planning_env):
    """
    Test C: Duplicate Planning ID
    Attempt to assign Planning ID 102 to two sections.
    Expected: Save is rejected with clear error message, existing intact.
    """
    client, db = client_and_db
    db.expire_all()
    s_print = db.query(Section).filter(Section.name == "PID Printing").first()
    s_cut = db.query(Section).filter(Section.name == "PID Cutting").first()

    # s_cut is Planning ID 101, s_print is Planning ID 102.
    # Attempt to change s_cut to Planning ID 102 without batch reorder
    payload = {
        "name": s_cut.name,
        "category_id": s_cut.category_id,
        "display_order": s_cut.display_order,
        "planning_id": 102,  # Conflicting with s_print
        "is_active": True,
        "counting_unit": "Pairs",
        "reason": "Attempt duplicate"
    }
    resp = client.put(f"/api/v1/settings/sections/{s_cut.id}", json=payload, headers=admin_headers)
    assert resp.status_code == 400
    err_detail = resp.json()["detail"]
    assert "Planning ID 102 is already assigned to another section" in err_detail

    # Verify database was NOT changed
    db.refresh(s_cut)
    assert s_cut.planning_id == 101


def test_d_section_rename_preserves_position_and_targets(client_and_db, admin_headers, footwear_planning_env):
    """
    Test D: Section Rename
    Rename Printing to Logo Printing.
    Expected: Table label changes, column position stays determined by Planning ID,
    target and production history remain correctly associated to internal ID.
    """
    client, db = client_and_db
    buyer = footwear_planning_env["buyer"]
    db.expire_all()
    s_print = db.query(Section).filter(Section.name == "PID Printing").first()

    rename_payload = {
        "name": "PID Logo Printing",
        "category_id": s_print.category_id,
        "display_order": s_print.display_order,
        "planning_id": s_print.planning_id,
        "is_active": True,
        "counting_unit": "Pairs",
        "minimum_daily_target": 1200,
        "maximum_daily_target": 3000,
        "reason": "Department branding update"
    }
    resp = client.put(f"/api/v1/settings/sections/{s_print.id}", json=rename_payload, headers=admin_headers)
    assert resp.status_code == 200
    db.expire_all()

    order = db.query(Order).filter(Order.order_code == "ORD-TEST-B-PID").first()
    sheet = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)

    sec_names = [s["section_name"] for s in sheet["sections"]]
    assert "PID LOGO PRINTING" in sec_names
    # Logo Printing still at Planning ID 102
    assert sheet["sections"][1]["section_name"] == "PID LOGO PRINTING"
    assert sheet["summary"]["PID LOGO PRINTING"]["effective_minimum_target"] == 1200
    assert sheet["summary"]["PID LOGO PRINTING"]["maximum_daily_target"] == 3000
    assert sheet["summary"]["PID LOGO PRINTING"]["total_ot"] == 480.0


def test_e_deactivated_section_excluded_without_breaking_gaps(client_and_db, admin_headers, footwear_planning_env):
    """
    Test E: Active Sections & Missing ID Gaps
    Deactivate Stockfitting (Planning ID 105) while keeping Cutting (101), Logo Printing (102), Lasting (103), Stitching (104).
    Expected: Deactivated section is excluded; remaining sections appear in ascending Planning ID order (101, 102, 103, 104).
    """
    client, db = client_and_db
    cats = footwear_planning_env["categories"]

    s_stock = db.query(Section).filter(Section.name == "PID Stockfitting").first()
    if not s_stock:
        s_stock = Section(name="PID Stockfitting", category_id=cats["fin"].id, planning_id=105, display_order=105, is_active=True)
        db.add(s_stock)
        db.commit()
        db.refresh(s_stock)

    # Deactivate Stockfitting
    resp = client.post(f"/api/v1/settings/sections/{s_stock.id}/deactivate", json={"force_confirm": True}, headers=admin_headers)
    assert resp.status_code == 200
    db.expire_all()

    order = db.query(Order).filter(Order.order_code == "ORD-TEST-B-PID").first()
    sheet = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)

    sec_names = [s["section_name"] for s in sheet["sections"]]
    assert "PID STOCKFITTING" not in sec_names
    assert len(sec_names) == 4
    # Gaps in ID sequence (if any) are smoothly accepted


def test_f_missing_target_settings_still_calculates_automatic_target(client_and_db, footwear_planning_env):
    """
    Test F: Missing Target Settings
    Leave Minimum and Maximum Target unset for an active section.
    Expected: Section still appears in correct Planning ID position,
    and automatic target engine calculates its Required Daily Rate without error.
    """
    _, db = client_and_db
    s_stitch = db.query(Section).filter(Section.name == "PID Stitching").first()
    # Zero out target setting
    ts = db.query(ProductionTargetSetting).filter(ProductionTargetSetting.section_id == s_stitch.id).first()
    if ts:
        ts.minimum_daily_target = 0
        ts.maximum_daily_target = 0
        db.commit()

    order = db.query(Order).filter(Order.order_code == "ORD-TEST-B-PID").first()
    sheet = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)

    stitch_summary = sheet["summary"]["PID STITCHING"]
    assert stitch_summary["planning_id"] == 104
    assert stitch_summary["required_daily_rate"] > 0
    assert stitch_summary["daily_target"] == stitch_summary["required_daily_rate"]


def test_g_wip_transfer_integrity_separated_from_display_order(client_and_db, footwear_planning_env):
    """
    Test G: WIP Integrity
    The editable Planning ID controls planning display order, but physical WIP flow
    continues to transfer along the order route workflow sequence.
    """
    _, db = client_and_db
    order = db.query(Order).filter(Order.order_code == "ORD-TEST-B-PID").first()

    # In ORD-TEST-B-PID, order routes sequence order:
    # 1: Cutting, 2: Printing, 3: Stitching, 4: Lasting
    # Display Planning ID:
    # 101: Cutting, 102: Logo Printing, 103: Lasting, 104: Stitching
    sheet = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)

    # Display columns: Cutting -> Logo Printing -> Lasting -> Stitching
    display_names = [s["section_name"] for s in sheet["sections"]]
    assert display_names == ["PID CUTTING", "PID LOGO PRINTING", "PID LASTING", "PID STITCHING"]

    # Physical workflow sequence remains untouched:
    workflow_routes = sorted(order.routes, key=lambda r: r.sequence_order)
    workflow_sec_names = [r.section.name for r in workflow_routes]
    assert workflow_sec_names == ["PID Cutting", "PID Logo Printing", "PID Stitching", "PID Lasting"]


def test_h_xlsx_and_csv_export_follow_planning_id_order(client_and_db, footwear_planning_env):
    """
    Test H: Export Consistency
    Excel and CSV exports follow ascending Planning ID order matching the UI.
    """
    _, db = client_and_db
    order = db.query(Order).filter(Order.order_code == "ORD-TEST-B-PID").first()

    # 1. XLSX
    xlsx_stream = ExportService.generate_matrix_planning_xlsx(db, order)
    wb = openpyxl.load_workbook(xlsx_stream)
    ws = wb["Production Planning"]
    # Row 4 has section group headers starting col 2, 5, 8, 11
    sec1 = ws.cell(row=4, column=2).value
    sec2 = ws.cell(row=4, column=5).value
    sec3 = ws.cell(row=4, column=8).value
    sec4 = ws.cell(row=4, column=11).value

    assert "PID CUTTING" in sec1
    assert "PID LOGO PRINTING" in sec2
    assert "PID LASTING" in sec3
    assert "PID STITCHING" in sec4

    # 2. CSV
    csv_str = ExportService.generate_matrix_planning_csv(db, order)
    reader = csv.reader(io.StringIO(csv_str))
    header = next(reader)
    # Header: Date, PID CUTTING T/G, PID CUTTING ON HAND, PID CUTTING O/T, PID LOGO PRINTING ..., PID LASTING ..., PID STITCHING ...
    assert header[1] == "PID CUTTING T/G"
    assert header[4] == "PID LOGO PRINTING T/G"
    assert header[7] == "PID LASTING T/G"
    assert header[10] == "PID STITCHING T/G"


def test_i_database_persistence_across_session(client_and_db, footwear_planning_env):
    """
    Test I: Persistence
    Restart session and ensure Planning IDs and ordering persist deterministically.
    """
    _, db = client_and_db
    db.expire_all()

    s_cut = db.query(Section).filter(Section.name == "PID Cutting").first()
    s_print = db.query(Section).filter(Section.name == "PID Logo Printing").first()
    s_last = db.query(Section).filter(Section.name == "PID Lasting").first()
    s_stitch = db.query(Section).filter(Section.name == "PID Stitching").first()

    assert s_cut.planning_id == 101
    assert s_print.planning_id == 102
    assert s_last.planning_id == 103
    assert s_stitch.planning_id == 104
