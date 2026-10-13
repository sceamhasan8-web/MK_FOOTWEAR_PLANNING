import io
import csv
from datetime import date, datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from app.models.order import Order
from app.models.route import OrderOperation
from app.models.entry import DailyProductionEntry
from app.services.order_service import OrderService
from app.services.calendar_service import CalendarService
from app.services.calculation_engine import CalculationEngine
from app.services.production_planning_service import ProductionPlanningService


class ExportService:

    @classmethod
    def generate_order_planning_xlsx(cls, db: Session, order: Order) -> io.BytesIO:
        """
        Generates a complete 7-sheet Order Planning Workbook.
        Sheets:
        1. Order Summary
        2. Daily Production Plan
        3. Section Summary
        4. Operation Details
        5. Component Details
        6. Production Entry History
        7. Delivery Date History
        """
        details = OrderService.get_order_calculated_details(db, order)
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Common styling tokens
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        def style_header_row(ws, row_idx, max_col):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

        def autofit_columns(ws):
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # -------------------------------------------------------------
        # SHEET 1: Order Summary
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="Order Summary")
        ws1["A1"] = "FOOTWEAR MANUFACTURING PRODUCTION ORDER SUMMARY"
        ws1["A1"].font = title_font

        summary_rows = [
            ("Order Code", details["order_code"]),
            ("Buyer", details["buyer_name"] or "N/A"),
            ("Customer Reference", details["customer_ref"]),
            ("Purchase Order (PO)", details["po_number"] or "N/A"),
            ("Product Style / Model", details["style_model"]),
            ("Description", details["description"] or "N/A"),
            ("Total Order Quantity (Pairs)", details["total_quantity"]),
            ("Original Order Date", str(details["original_order_date"])),
            ("Production Start Date", str(details["production_start_date"])),
            ("Original Delivery Date", str(details["original_delivery_date"])),
            ("Current Delivery Date", str(details["current_delivery_date"])),
            ("Order Priority", str(details["priority"].value if hasattr(details["priority"], "value") else details["priority"])),
            ("Overall Status", str(details["status"].value if hasattr(details["status"], "value") else details["status"])),
            ("Remaining Working Days", details["remaining_working_days"]),
            ("Completed Pairs", details["completed_pairs"]),
            ("Outstanding Pairs", details["outstanding_pairs"]),
            ("Completion Percentage", f"{details['completion_percentage']}%"),
            ("Is Overdue", "YES" if details["is_overdue"] else "NO"),
            ("Is At Risk", "YES" if details["is_at_risk"] else "NO"),
        ]

        for i, (k, v) in enumerate(summary_rows, start=3):
            c_k = ws1.cell(row=i, column=1, value=k)
            c_k.font = Font(bold=True)
            c_k.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
            c_v = ws1.cell(row=i, column=2, value=v)
            c_k.border = thin_border
            c_v.border = thin_border

        autofit_columns(ws1)

        # -------------------------------------------------------------
        # SHEET 2: Daily Production Plan
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Daily Production Plan")
        ws2["A1"] = f"DAILY PRODUCTION PLAN: {details['order_code']}"
        ws2["A1"].font = title_font

        headers_plan = [
            "Date", "Day", "Working Day Status", "Operation", 
            "Opening Balance", "Planned Qty", "Daily Target", "Cumulative Plan"
        ]
        ws2.row_dimensions[3].height = 25
        for col_idx, h in enumerate(headers_plan, start=1):
            ws2.cell(row=3, column=col_idx, value=h)
        style_header_row(ws2, 3, len(headers_plan))

        curr_row = 4
        for r in details["routes"]:
            if r["status"] == "REQUIRED":
                for op in r["operations"]:
                    op_id = op["id"]
                    op_obj = db.query(OrderOperation).filter(OrderOperation.id == op_id).first()
                    if op_obj:
                        future_plan = ProductionPlanningService.generate_future_daily_plan(db, order, op_obj)
                        cum_plan = 0.0
                        for row_p in future_plan:
                            p_qty = float(row_p.get("auto_target_numeric", 0.0))
                            cum_plan += p_qty
                            ws2.cell(row=curr_row, column=1, value=str(row_p["date"])).border = thin_border
                            ws2.cell(row=curr_row, column=2, value=row_p["day_name"]).border = thin_border
                            ws2.cell(row=curr_row, column=3, value="Working Day").border = thin_border
                            ws2.cell(row=curr_row, column=4, value=op["operation_name"]).border = thin_border
                            ws2.cell(row=curr_row, column=5, value=row_p["opening_balance"]).border = thin_border
                            ws2.cell(row=curr_row, column=6, value=p_qty).border = thin_border
                            ws2.cell(row=curr_row, column=7, value=op["daily_target"]).border = thin_border
                            ws2.cell(row=curr_row, column=8, value=cum_plan).border = thin_border
                            curr_row += 1

        autofit_columns(ws2)

        # -------------------------------------------------------------
        # SHEET 3: Section Summary
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Section Summary")
        ws3["A1"] = f"SECTION-WISE PRODUCTION SUMMARY: {details['order_code']}"
        ws3["A1"].font = title_font

        headers_sec = [
            "Section Name", "Category", "Routing Status", 
            "Operations Count", "Section Complete?", "Component Tracking"
        ]
        for col_idx, h in enumerate(headers_sec, start=1):
            ws3.cell(row=3, column=col_idx, value=h)
        style_header_row(ws3, 3, len(headers_sec))

        for idx, r in enumerate(details["routes"], start=4):
            ws3.cell(row=idx, column=1, value=r["section_name"]).border = thin_border
            ws3.cell(row=idx, column=2, value=r["category_name"]).border = thin_border
            ws3.cell(row=idx, column=3, value=r["status"].value if hasattr(r["status"], "value") else str(r["status"])).border = thin_border
            ws3.cell(row=idx, column=4, value=len(r["operations"])).border = thin_border
            ws3.cell(row=idx, column=5, value="YES" if r["is_completed"] else "NO").border = thin_border
            ws3.cell(row=idx, column=6, value="YES" if r["has_components"] else "NO").border = thin_border

        autofit_columns(ws3)

        # -------------------------------------------------------------
        # SHEET 4: Operation Details
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="Operation Details")
        ws4["A1"] = f"OPERATION-LEVEL DETAILS: {details['order_code']}"
        ws4["A1"].font = title_font

        headers_ops = [
            "Section", "Operation Name", "Unit", "Required Qty", 
            "Cumulative Accepted", "Cumulative Rejected", "Outstanding Balance", 
            "Completion %", "Daily Target", "Operation Status"
        ]
        for col_idx, h in enumerate(headers_ops, start=1):
            ws4.cell(row=3, column=col_idx, value=h)
        style_header_row(ws4, 3, len(headers_ops))

        op_row = 4
        for r in details["routes"]:
            for op in r["operations"]:
                ws4.cell(row=op_row, column=1, value=r["section_name"]).border = thin_border
                ws4.cell(row=op_row, column=2, value=op["operation_name"]).border = thin_border
                ws4.cell(row=op_row, column=3, value=op["unit"].value if hasattr(op["unit"], "value") else str(op["unit"])).border = thin_border
                ws4.cell(row=op_row, column=4, value=op["required_quantity"]).border = thin_border
                ws4.cell(row=op_row, column=5, value=op["completed_quantity"]).border = thin_border
                ws4.cell(row=op_row, column=6, value=op["rejected_quantity"]).border = thin_border
                ws4.cell(row=op_row, column=7, value=op["balance_quantity"]).border = thin_border
                ws4.cell(row=op_row, column=8, value=f"{op['completion_percentage']}%").border = thin_border
                ws4.cell(row=op_row, column=9, value=op["daily_target"]).border = thin_border
                ws4.cell(row=op_row, column=10, value=op["status"].value if hasattr(op["status"], "value") else str(op["status"])).border = thin_border
                op_row += 1

        autofit_columns(ws4)

        # -------------------------------------------------------------
        # SHEET 5: Component Details
        # -------------------------------------------------------------
        ws5 = wb.create_sheet(title="Component Details")
        ws5["A1"] = f"COMPONENT-WISE PRODUCTION: {details['order_code']}"
        ws5["A1"].font = title_font

        headers_comp = [
            "Section", "Operation", "Component Name", "Unit", 
            "Pieces/Pair", "Required Total", "Accepted Total", 
            "Rejected Total", "Balance Total", "Pair-Equivalent Accepted"
        ]
        for col_idx, h in enumerate(headers_comp, start=1):
            ws5.cell(row=3, column=col_idx, value=h)
        style_header_row(ws5, 3, len(headers_comp))

        comp_row = 4
        for r in details["routes"]:
            for op in r["operations"]:
                for c in op["components"]:
                    ws5.cell(row=comp_row, column=1, value=r["section_name"]).border = thin_border
                    ws5.cell(row=comp_row, column=2, value=op["operation_name"]).border = thin_border
                    ws5.cell(row=comp_row, column=3, value=c["component_name"]).border = thin_border
                    ws5.cell(row=comp_row, column=4, value=c["unit"].value if hasattr(c["unit"], "value") else str(c["unit"])).border = thin_border
                    ws5.cell(row=comp_row, column=5, value=c["pieces_per_pair"]).border = thin_border
                    ws5.cell(row=comp_row, column=6, value=c["required_total_qty"]).border = thin_border
                    ws5.cell(row=comp_row, column=7, value=c["accepted_total_qty"]).border = thin_border
                    ws5.cell(row=comp_row, column=8, value=c["rejected_total_qty"]).border = thin_border
                    ws5.cell(row=comp_row, column=9, value=c["balance_total_qty"]).border = thin_border
                    ws5.cell(row=comp_row, column=10, value=c["pair_equivalent_accepted"]).border = thin_border
                    comp_row += 1

        autofit_columns(ws5)

        # -------------------------------------------------------------
        # SHEET 6: Production Entry History
        # -------------------------------------------------------------
        ws6 = wb.create_sheet(title="Production Entry History")
        ws6["A1"] = f"DAILY PRODUCTION ENTRIES: {details['order_code']}"
        ws6["A1"].font = title_font

        headers_hist = [
            "Entry ID", "Entry Date", "Section", "Operation", "Component", 
            "Quantity", "Accepted Qty", "Rejected Qty", "Unit", "Entered By", "Reversed?", "Remarks"
        ]
        for col_idx, h in enumerate(headers_hist, start=1):
            ws6.cell(row=3, column=col_idx, value=h)
        style_header_row(ws6, 3, len(headers_hist))

        entries = db.query(DailyProductionEntry).filter(
            DailyProductionEntry.order_id == order.id
        ).order_by(DailyProductionEntry.entry_date.desc(), DailyProductionEntry.id.desc()).all()

        for idx, e in enumerate(entries, start=4):
            ws6.cell(row=idx, column=1, value=e.id).border = thin_border
            ws6.cell(row=idx, column=2, value=str(e.entry_date)).border = thin_border
            ws6.cell(row=idx, column=3, value=e.route.section.name if (e.route and e.route.section) else "").border = thin_border
            ws6.cell(row=idx, column=4, value=e.operation.operation_name if e.operation else "").border = thin_border
            ws6.cell(row=idx, column=5, value=e.component.component_name if e.component else "N/A").border = thin_border
            ws6.cell(row=idx, column=6, value=e.quantity).border = thin_border
            ws6.cell(row=idx, column=7, value=e.accepted_quantity).border = thin_border
            ws6.cell(row=idx, column=8, value=e.rejected_quantity).border = thin_border
            ws6.cell(row=idx, column=9, value=e.unit).border = thin_border
            ws6.cell(row=idx, column=10, value=e.entered_by_user.full_name if e.entered_by_user else "").border = thin_border
            ws6.cell(row=idx, column=11, value="YES" if e.is_reversed else "NO").border = thin_border
            ws6.cell(row=idx, column=12, value=e.remarks or "").border = thin_border

        autofit_columns(ws6)

        # -------------------------------------------------------------
        # SHEET 7: Delivery Date History
        # -------------------------------------------------------------
        ws7 = wb.create_sheet(title="Delivery Date History")
        ws7["A1"] = f"DELIVERY DATE REVISION HISTORY: {details['order_code']}"
        ws7["A1"].font = title_font

        headers_rev = [
            "Revision ID", "Previous Delivery Date", "Revised Delivery Date", 
            "Reason", "Revised By", "Revision Timestamp"
        ]
        for col_idx, h in enumerate(headers_rev, start=1):
            ws7.cell(row=3, column=col_idx, value=h)
        style_header_row(ws7, 3, len(headers_rev))

        for idx, rev in enumerate(details["delivery_revisions"], start=4):
            ws7.cell(row=idx, column=1, value=rev["id"]).border = thin_border
            ws7.cell(row=idx, column=2, value=str(rev["previous_delivery_date"])).border = thin_border
            ws7.cell(row=idx, column=3, value=str(rev["revised_delivery_date"])).border = thin_border
            ws7.cell(row=idx, column=4, value=rev["reason"]).border = thin_border
            ws7.cell(row=idx, column=5, value=rev["revised_by_name"] or "").border = thin_border
            ws7.cell(row=idx, column=6, value=str(rev["created_at"])).border = thin_border

        autofit_columns(ws7)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def generate_order_csv(cls, db: Session, order: Order) -> str:
        """Exports order operation details to CSV."""
        details = OrderService.get_order_calculated_details(db, order)
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(["Order Code", details["order_code"]])
        writer.writerow(["Buyer", details["buyer_name"]])
        writer.writerow(["Style Model", details["style_model"]])
        writer.writerow(["Total Quantity", details["total_quantity"]])
        writer.writerow(["Current Delivery Date", details["current_delivery_date"]])
        writer.writerow(["Status", details["status"].value if hasattr(details["status"], "value") else str(details["status"])])
        writer.writerow([])

        writer.writerow([
            "Section", "Operation", "Unit", "Required Qty", 
            "Cumulative Accepted", "Cumulative Rejected", "Balance", 
            "Completion Pct", "Daily Target", "Status"
        ])

        for r in details["routes"]:
            for op in r["operations"]:
                writer.writerow([
                    r["section_name"],
                    op["operation_name"],
                    op["unit"].value if hasattr(op["unit"], "value") else str(op["unit"]),
                    op["required_quantity"],
                    op["completed_quantity"],
                    op["rejected_quantity"],
                    op["balance_quantity"],
                    f"{op['completion_percentage']}%",
                    op["daily_target"],
                    op["status"].value if hasattr(op["status"], "value") else str(op["status"]),
                ])

        return output.getvalue()

    @classmethod
    def generate_category_flow_planning_xlsx(
        cls,
        db: Session,
        order_id: Optional[int] = None
    ) -> io.BytesIO:
        """
        Generates official 4-sheet Production_Planning.xlsx workbook (Section 32):
        Sheet 1: Master Planning
        Sheet 2: Order Summary
        Sheet 3: Daily Summary
        Sheet 4: Category Flow
        """
        rows = ProductionPlanningService.generate_category_flow_planning_sheet(db, order_id=order_id)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default blank sheet

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        def style_header_row(ws, row_idx, max_col):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        def autofit(ws):
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # -------------------------------------------------------------
        # SHEET 1: Master Planning
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="Master Planning")
        ws1["A1"] = "AUTOMATIC PRODUCTION FLOW MASTER PLANNING SHEET"
        ws1["A1"].font = title_font

        headers1 = [
            "Date", "Buyer", "Order", "Order Qty", "Category", "Section", "Process",
            "Opening On Hand", "Incoming WIP", "Available On Hand", "Daily Target",
            "Executable Target", "Actual", "Output Completed", "Closing On Hand",
            "Balance", "Variance", "WIP Status", "Production Status", "Shipment Date"
        ]
        ws1.row_dimensions[3].height = 25
        for col_idx, h in enumerate(headers1, start=1):
            ws1.cell(row=3, column=col_idx, value=h)
        style_header_row(ws1, 3, len(headers1))

        for r_idx, r in enumerate(rows, start=4):
            data = [
                str(r["date"]),
                r["buyer_name"],
                r["order_code"],
                r["order_quantity"],
                r["category_name"],
                r["section_name"],
                r["operation_name"],
                r["opening_on_hand"],
                r["incoming_wip"],
                r["available_on_hand"],
                r["daily_target"],
                r["executable_target"],
                r["actual_production"] if r["actual_production"] is not None else "-",
                r["completed_output"],
                r["closing_on_hand"],
                r["balance"],
                r["variance"] if r["variance"] is not None else "-",
                r["wip_status"],
                r["production_status"],
                str(r["shipment_date"]),
            ]
            for c_idx, val in enumerate(data, start=1):
                cell = ws1.cell(row=r_idx, column=c_idx, value=val)
                cell.border = thin_border
                if c_idx in [4, 8, 9, 10, 11, 12, 14, 15, 16]:
                    cell.number_format = "#,##0"

        ws1.freeze_panes = "A4"
        autofit(ws1)

        # -------------------------------------------------------------
        # SHEET 2: Order Summary
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Order Summary")
        ws2["A1"] = "ORDER PRODUCTION FLOW SUMMARY"
        ws2["A1"].font = title_font

        headers2 = [
            "Buyer", "Order", "Order Qty", "Category 1 Completed", "Category 2 Completed",
            "Category 3 Completed", "Final Completed Qty", "Final Balance", "Shipment Date",
            "Days Remaining", "Overall Status"
        ]
        ws2.row_dimensions[3].height = 25
        for col_idx, h in enumerate(headers2, start=1):
            ws2.cell(row=3, column=col_idx, value=h)
        style_header_row(ws2, 3, len(headers2))

        # Query orders for summary
        query_orders = db.query(Order)
        if order_id:
            query_orders = query_orders.filter(Order.id == order_id)
        orders_list = query_orders.order_by(Order.current_delivery_date.asc()).all()

        curr_r2 = 4
        for o in orders_list:
            flow = ProductionPlanningService.calculate_category_flow(db, o)
            plan_det = OrderService.get_order_calculated_details(db, o)

            c1_out = flow[0]["completed_output"] if len(flow) > 0 else 0.0
            c2_out = flow[1]["completed_output"] if len(flow) > 1 else 0.0
            c3_out = flow[2]["completed_output"] if len(flow) > 2 else 0.0
            final_out = flow[-1]["completed_output"] if flow else 0.0
            final_bal = max(0.0, o.total_quantity - final_out)

            row_data = [
                o.buyer.name if o.buyer else "N/A",
                o.order_code,
                o.total_quantity,
                c1_out,
                c2_out,
                c3_out,
                final_out,
                final_bal,
                str(o.current_delivery_date),
                plan_det.get("remaining_days", 0),
                str(plan_det.get("status", "")),
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws2.cell(row=curr_r2, column=c_idx, value=val)
                cell.border = thin_border
                if c_idx in [3, 4, 5, 6, 7, 8]:
                    cell.number_format = "#,##0"
            curr_r2 += 1

        ws2.freeze_panes = "A4"
        autofit(ws2)

        # -------------------------------------------------------------
        # SHEET 3: Daily Summary
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Daily Summary")
        ws3["A1"] = "DAILY FACTORY PRODUCTION & WIP SUMMARY"
        ws3["A1"].font = title_font

        headers3 = [
            "Date", "Total Target", "Total Actual", "Total WIP Incoming",
            "Total WIP Available", "Total Balance", "Orders At Risk", "Orders Overdue"
        ]
        ws3.row_dimensions[3].height = 25
        for col_idx, h in enumerate(headers3, start=1):
            ws3.cell(row=3, column=col_idx, value=h)
        style_header_row(ws3, 3, len(headers3))

        # Aggregate rows by date
        dates_summary: Dict[date, Dict[str, float]] = {}
        for r in rows:
            d = r["date"]
            if d not in dates_summary:
                dates_summary[d] = {
                    "target": 0.0,
                    "actual": 0.0,
                    "incoming": 0.0,
                    "available": 0.0,
                    "balance": 0.0,
                }
            dates_summary[d]["target"] += r["daily_target"]
            if r["actual_production"] is not None:
                dates_summary[d]["actual"] += r["actual_production"]
            dates_summary[d]["incoming"] += r["incoming_wip"]
            dates_summary[d]["available"] += r["available_on_hand"]
            dates_summary[d]["balance"] += r["balance"]

        curr_r3 = 4
        for d in sorted(dates_summary.keys()):
            ds = dates_summary[d]
            row_data = [
                str(d),
                ds["target"],
                ds["actual"],
                ds["incoming"],
                ds["available"],
                ds["balance"],
                sum(1 for o in orders_list if o.current_delivery_date >= d and ProductionPlanningService.calculate_order_plan(db, o).get("is_at_risk")),
                sum(1 for o in orders_list if o.current_delivery_date < d and ProductionPlanningService.calculate_order_plan(db, o).get("balance_quantity", 0) > 0),
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws3.cell(row=curr_r3, column=c_idx, value=val)
                cell.border = thin_border
                if c_idx in [2, 3, 4, 5, 6]:
                    cell.number_format = "#,##0"
            curr_r3 += 1

        ws3.freeze_panes = "A4"
        autofit(ws3)

        # -------------------------------------------------------------
        # SHEET 4: Category Flow
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="Category Flow")
        ws4["A1"] = "CATEGORY-TO-CATEGORY WIP FLOW SHEET"
        ws4["A1"].font = title_font

        headers4 = [
            "Date", "Category 1 Output", "Category 2 Incoming",
            "Category 2 Output", "Category 3 Incoming", "Category 3 Output"
        ]
        ws4.row_dimensions[3].height = 25
        for col_idx, h in enumerate(headers4, start=1):
            ws4.cell(row=3, column=col_idx, value=h)
        style_header_row(ws4, 3, len(headers4))

        # Build category flow by date
        cat_flow_dates: Dict[date, Dict[str, float]] = {}
        for r in rows:
            d = r["date"]
            if d not in cat_flow_dates:
                cat_flow_dates[d] = {
                    "c1_out": 0.0, "c2_in": 0.0, "c2_out": 0.0,
                    "c3_in": 0.0, "c3_out": 0.0
                }
            cat_name = str(r["category_name"]).lower()
            if "cutting" in cat_name or "category 1" in cat_name:
                cat_flow_dates[d]["c1_out"] += r["completed_output"]
            elif "preparation" in cat_name or "category 2" in cat_name:
                cat_flow_dates[d]["c2_in"] += r["incoming_wip"]
                cat_flow_dates[d]["c2_out"] += r["completed_output"]
            elif "lasting" in cat_name or "dip" in cat_name or "category 3" in cat_name:
                cat_flow_dates[d]["c3_in"] += r["incoming_wip"]
                cat_flow_dates[d]["c3_out"] += r["completed_output"]

        curr_r4 = 4
        for d in sorted(cat_flow_dates.keys()):
            cf = cat_flow_dates[d]
            row_data = [
                str(d),
                cf["c1_out"],
                cf["c2_in"],
                cf["c2_out"],
                cf["c3_in"],
                cf["c3_out"],
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws4.cell(row=curr_r4, column=c_idx, value=val)
                cell.border = thin_border
                if c_idx >= 2:
                    cell.number_format = "#,##0"
            curr_r4 += 1

        ws4.freeze_panes = "A4"
        autofit(ws4)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    def generate_matrix_planning_xlsx(cls, db: Session, order: Order) -> io.BytesIO:
        """
        Generates the 5-sheet Date-Wise TG / ON HAND / O/T Planning Workbook (Sections 39, 40).
        Sheet 1: Production Planning (Date x Section Matrix)
        Sheet 2: Summary (Section | Total T/G | Final On Hand | Total O/T | Balance | Status)
        Sheet 3: Order Details
        Sheet 4: Production Flow
        Sheet 5: Daily Actual
        """
        matrix = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default blank sheet

        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        section_header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        sub_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        summary_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        title_font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        double_bottom_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='double', color='1F4E79')
        )

        def autofit(ws):
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        # -------------------------------------------------------------
        # SHEET 1: Production Planning
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="Production Planning")
        ws1["A1"] = f"DATE-WISE PRODUCTION PLANNING SHEET — {matrix['order_code']}"
        ws1["A1"].font = title_font

        ws1["A2"] = f"Buyer: {matrix['buyer_name']} | Order Qty: {matrix['total_quantity']:,.0f} | Shipment Date: {matrix['shipment_date']} | Days Remaining: {matrix['remaining_working_days']} | Status: {matrix['overall_status']}"
        ws1["A2"].font = Font(name="Calibri", size=10, italic=True)

        sections = matrix.get("sections", [])
        rows = matrix.get("rows", [])
        summary = matrix.get("summary", {})

        # Row 4: Section group headers
        # Row 5: Date, then T/G, ON HAND, O/T for each section
        ws1.row_dimensions[4].height = 22
        ws1.row_dimensions[5].height = 20

        # Col 1: Date
        c_date = ws1.cell(row=4, column=1, value="Date")
        c_date.font = header_font
        c_date.fill = section_header_fill
        c_date.alignment = Alignment(horizontal="center", vertical="center")
        c_date.border = thin_border
        ws1.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)

        col_idx = 2
        for s in sections:
            s_name = s["section_name"]
            # Merge 3 columns for section name in row 4
            ws1.merge_cells(start_row=4, start_column=col_idx, end_row=4, end_column=col_idx + 2)
            c_sec = ws1.cell(row=4, column=col_idx, value=f"{s_name} ({'REQ' if s['is_required'] else 'N/A'})")
            c_sec.font = header_font
            c_sec.fill = section_header_fill
            c_sec.alignment = Alignment(horizontal="center", vertical="center")
            c_sec.border = thin_border

            # Sub-headers in row 5
            for offset, label in enumerate(["T/G", "ON HAND", "O/T"]):
                c_sub = ws1.cell(row=5, column=col_idx + offset, value=f"{s_name} {label}")
                c_sub.font = header_font
                c_sub.fill = sub_header_fill
                c_sub.alignment = Alignment(horizontal="center", vertical="center")
                c_sub.border = thin_border

            col_idx += 3

        # Populate date rows starting at row 6
        curr_row = 6
        for r in rows:
            ws1.cell(row=curr_row, column=1, value=str(r["date"])).border = thin_border
            ws1.cell(row=curr_row, column=1).alignment = Alignment(horizontal="center")

            c_pos = 2
            for s in sections:
                s_name = s["section_name"]
                s_data = r["section_data"].get(s_name, {})
                tg_val = s_data.get("tg", 0.0)
                oh_val = s_data.get("on_hand", 0.0)
                ot_val = s_data.get("ot", 0.0)

                c_tg = ws1.cell(row=curr_row, column=c_pos, value=tg_val)
                c_oh = ws1.cell(row=curr_row, column=c_pos + 1, value=oh_val)
                c_ot = ws1.cell(row=curr_row, column=c_pos + 2, value=ot_val)

                for c in [c_tg, c_oh, c_ot]:
                    c.border = thin_border
                    c.number_format = "#,##0"

                c_pos += 3
            curr_row += 1

        # Summary Row (Row curr_row)
        ws1.cell(row=curr_row, column=1, value="TOTAL / SUMMARY").border = double_bottom_border
        ws1.cell(row=curr_row, column=1).font = bold_font
        ws1.cell(row=curr_row, column=1).fill = summary_fill
        ws1.cell(row=curr_row, column=1).alignment = Alignment(horizontal="center")

        c_pos = 2
        for s in sections:
            s_name = s["section_name"]
            s_sum = summary.get(s_name, {})
            tot_tg = s_sum.get("total_tg", 0.0)
            fin_oh = s_sum.get("final_on_hand", 0.0)
            tot_ot = s_sum.get("total_ot", 0.0)

            c_tg = ws1.cell(row=curr_row, column=c_pos, value=tot_tg)
            c_oh = ws1.cell(row=curr_row, column=c_pos + 1, value=fin_oh)
            c_ot = ws1.cell(row=curr_row, column=c_pos + 2, value=tot_ot)

            for c in [c_tg, c_oh, c_ot]:
                c.border = double_bottom_border
                c.font = bold_font
                c.fill = summary_fill
                c.number_format = "#,##0"

            c_pos += 3

        ws1.freeze_panes = "B6"
        autofit(ws1)

        # -------------------------------------------------------------
        # SHEET 2: Summary
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Summary")
        ws2["A1"] = "SECTION PRODUCTION PLANNING SUMMARY"
        ws2["A1"].font = title_font

        headers2 = ["Section", "Total T/G", "Final On Hand", "Total O/T", "Balance", "Status"]
        ws2.row_dimensions[3].height = 22
        for col_idx, h in enumerate(headers2, start=1):
            cell = ws2.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = section_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        curr_r2 = 4
        for s in sections:
            s_name = s["section_name"]
            s_sum = summary.get(s_name, {})
            row_vals = [
                s_name,
                s_sum.get("total_tg", 0.0),
                s_sum.get("final_on_hand", 0.0),
                s_sum.get("total_ot", 0.0),
                s_sum.get("balance", 0.0),
                s_sum.get("status", "N/A")
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws2.cell(row=curr_r2, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in [2, 3, 4, 5]:
                    cell.number_format = "#,##0"
            curr_r2 += 1

        autofit(ws2)

        # -------------------------------------------------------------
        # SHEET 3: Order Details
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Order Details")
        ws3["A1"] = f"ORDER SPECIFICATIONS — {order.order_code}"
        ws3["A1"].font = title_font

        details_list = [
            ("Order Code", order.order_code),
            ("Buyer", order.buyer.name if order.buyer else "N/A"),
            ("Customer Reference", order.customer_ref or "N/A"),
            ("PO Number", order.po_number or "N/A"),
            ("Style / Model", order.style_model or "N/A"),
            ("Total Quantity (Pairs)", order.total_quantity),
            ("Production Start Date", str(order.production_start_date)),
            ("Shipment / Delivery Date", str(order.current_delivery_date)),
            ("Remaining Working Days", matrix.get("remaining_working_days", 0)),
            ("Final Balance (Pairs)", matrix.get("final_balance", 0.0)),
            ("Overall Status", matrix.get("overall_status", "N/A")),
        ]
        for i, (k, v) in enumerate(details_list, start=3):
            ck = ws3.cell(row=i, column=1, value=k)
            cv = ws3.cell(row=i, column=2, value=v)
            ck.font = bold_font
            ck.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
            ck.border = thin_border
            cv.border = thin_border
            if isinstance(v, (int, float)):
                cv.number_format = "#,##0"

        autofit(ws3)

        # -------------------------------------------------------------
        # SHEET 4: Production Flow
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="Production Flow")
        ws4["A1"] = "CATEGORY-TO-CATEGORY ON-HAND WIP FLOW SUMMARY"
        ws4["A1"].font = title_font

        headers4 = [
            "Category", "Required Quantity", "Opening On Hand", "Incoming WIP",
            "Available On Hand", "Actual Consumed", "Completed Output", "Closing On Hand",
            "Balance", "WIP Status"
        ]
        ws4.row_dimensions[3].height = 22
        for col_idx, h in enumerate(headers4, start=1):
            cell = ws4.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = section_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        curr_r4 = 4
        for cat in matrix.get("category_flow_summary", []):
            cat_vals = [
                cat.category_name if hasattr(cat, 'category_name') else cat.get('category_name'),
                cat.required_quantity if hasattr(cat, 'required_quantity') else cat.get('required_quantity'),
                cat.opening_on_hand if hasattr(cat, 'opening_on_hand') else cat.get('opening_on_hand'),
                cat.incoming_wip if hasattr(cat, 'incoming_wip') else cat.get('incoming_wip'),
                cat.available_on_hand if hasattr(cat, 'available_on_hand') else cat.get('available_on_hand'),
                cat.actual_production if hasattr(cat, 'actual_production') else cat.get('actual_production'),
                cat.completed_output if hasattr(cat, 'completed_output') else cat.get('completed_output'),
                cat.closing_on_hand if hasattr(cat, 'closing_on_hand') else cat.get('closing_on_hand'),
                cat.balance if hasattr(cat, 'balance') else cat.get('balance'),
                cat.wip_status if hasattr(cat, 'wip_status') else cat.get('wip_status'),
            ]
            for col_idx, val in enumerate(cat_vals, start=1):
                cell = ws4.cell(row=curr_r4, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in [2, 3, 4, 5, 6, 7, 8, 9]:
                    cell.number_format = "#,##0"
            curr_r4 += 1

        autofit(ws4)

        # -------------------------------------------------------------
        # SHEET 5: Daily Actual
        # -------------------------------------------------------------
        ws5 = wb.create_sheet(title="Daily Actual")
        ws5["A1"] = f"DAILY ACTUAL PRODUCTION ENTRIES — {order.order_code}"
        ws5["A1"].font = title_font

        headers5 = ["Date", "Route / Section", "Accepted Quantity", "Rejected", "Created At", "Remarks"]
        ws5.row_dimensions[3].height = 22
        for col_idx, h in enumerate(headers5, start=1):
            cell = ws5.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = section_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        all_entries = db.query(DailyProductionEntry).filter(
            DailyProductionEntry.order_id == order.id,
            DailyProductionEntry.is_reversed == False
        ).order_by(DailyProductionEntry.entry_date.asc()).all()

        curr_r5 = 4
        for e in all_entries:
            sec_name = e.route.section.name if (e.route and e.route.section) else "N/A"
            entry_vals = [
                str(e.entry_date),
                sec_name,
                e.accepted_quantity,
                e.rejected_quantity,
                str(e.created_at.strftime("%Y-%m-%d %H:%M") if e.created_at else ""),
                e.remarks or ""
            ]
            for col_idx, val in enumerate(entry_vals, start=1):
                cell = ws5.cell(row=curr_r5, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in [3, 4]:
                    cell.number_format = "#,##0"
            curr_r5 += 1

        autofit(ws5)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    def generate_matrix_planning_csv(cls, db: Session, order: Order) -> str:
        """
        Generates CSV format of the matrix planning table.
        """
        matrix = ProductionPlanningService.generate_order_matrix_planning_sheet(db, order.id)
        sections = matrix.get("sections", [])
        rows = matrix.get("rows", [])
        summary = matrix.get("summary", {})

        output = io.StringIO()
        writer = csv.writer(output)

        # Header 1: Date, then section names repeated 3 times
        h1 = ["Date"]
        for s in sections:
            s_name = s["section_name"]
            h1.extend([f"{s_name} T/G", f"{s_name} ON HAND", f"{s_name} O/T"])
        writer.writerow(h1)

        # Data rows
        for r in rows:
            row_data = [str(r["date"])]
            for s in sections:
                s_name = s["section_name"]
                s_data = r["section_data"].get(s_name, {})
                row_data.extend([
                    s_data.get("tg", 0.0),
                    s_data.get("on_hand", 0.0),
                    s_data.get("ot", 0.0)
                ])
            writer.writerow(row_data)

        # Summary row
        sum_row = ["TOTAL / SUMMARY"]
        for s in sections:
            s_name = s["section_name"]
            s_sum = summary.get(s_name, {})
            sum_row.extend([
                s_sum.get("total_tg", 0.0),
                s_sum.get("final_on_hand", 0.0),
                s_sum.get("total_ot", 0.0)
            ])
        writer.writerow(sum_row)

        return output.getvalue()

    @classmethod
    def generate_wip_flow_xlsx(cls, db: Session, order: Order) -> io.BytesIO:
        """
        Generates the authoritative 5-sheet Serial WIP Flow Planning Workbook (Sections 43 & 44).
        Sheets:
        1. WIP Flow Plan
        2. Workflow Summary
        3. Factory Calendar
        4. Production Entries
        5. Order Details
        """
        from app.services.wip_planning_engine import WIPPlanningEngine
        plan_data = WIPPlanningEngine.calculate_order_wip_flow(db, order)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Styles
        title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        summary_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        status_fills = {
            "COMPLETED": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
            "ON TRACK": PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid"),
            "AHEAD": PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid"),
            "WIP SHORTAGE": PatternFill(start_color="FFE8CC", end_color="FFE8CC", fill_type="solid"),
            "BEHIND TARGET": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
            "OVERDUE": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
            "AT RISK": PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid"),
        }

        def autofit(ws):
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # -------------------------------------------------------------
        # SHEET 1: WIP Flow Plan
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="WIP Flow Plan")
        ws1["A1"] = f"FOOTWEAR MANUFACTURING WIP FLOW PLANNING SHEET — ORDER: {order.order_code}"
        ws1["A1"].font = title_font
        ws1["A2"] = (
            f"Buyer: {plan_data['buyer_name']} | Quantity: {order.total_quantity:,.0f} Pairs | "
            f"Shipment: {order.current_delivery_date} | Bottleneck: {plan_data['current_bottleneck']} ({plan_data['wip_waiting']:,.0f} WIP Waiting)"
        )
        ws1["A2"].font = subtitle_font

        headers1 = [
            "Date", "Serial", "Section", "T/G", "Opening On Hand", "Incoming WIP",
            "Actual", "Complete", "Closing On Hand", "Balance", "Remaining Working Days", "Status"
        ]
        ws1.row_dimensions[4].height = 24
        for col_idx, h in enumerate(headers1, start=1):
            cell = ws1.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        curr_r1 = 5
        daily_rows = plan_data.get("daily_plan_rows", [])
        for row_dict in daily_rows:
            row_vals = [
                str(row_dict["date"]),
                row_dict["serial_no"],
                row_dict["section_name"],
                row_dict["target"],
                row_dict["opening_on_hand"],
                row_dict["incoming_wip"],
                row_dict["actual"],
                row_dict["complete"],
                row_dict["closing_on_hand"],
                row_dict["balance"],
                row_dict["remaining_working_days"],
                row_dict["status"]
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws1.cell(row=curr_r1, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in [4, 5, 6, 7, 8, 9, 10]:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in [1, 2, 11]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 12:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    stat_key = str(val).upper()
                    if stat_key in status_fills:
                        cell.fill = status_fills[stat_key]
            curr_r1 += 1

        # Summary Row
        ws1.row_dimensions[curr_r1].height = 22
        sum_cells = [
            "TOTAL / SUMMARY", "", "",
            sum(r["target"] for r in daily_rows),
            "", "",
            sum(r["actual"] for r in daily_rows),
            "", "", "", "", ""
        ]
        for col_idx, val in enumerate(sum_cells, start=1):
            cell = ws1.cell(row=curr_r1, column=col_idx, value=val)
            cell.font = Font(bold=True)
            cell.fill = summary_fill
            cell.border = thin_border
            if col_idx in [4, 7]:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")

        ws1.freeze_panes = "A5"
        ws1.auto_filter.ref = f"A4:L{curr_r1 - 1}"
        autofit(ws1)

        # -------------------------------------------------------------
        # SHEET 2: Workflow Summary
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Workflow Summary")
        ws2["A1"] = f"WORKFLOW PRODUCTION SUMMARY — {order.order_code}"
        ws2["A1"].font = title_font

        headers2 = ["S/N", "Section", "Required Qty", "Total T/G", "Complete", "Final On Hand", "Balance", "Status"]
        ws2.row_dimensions[3].height = 22
        for col_idx, h in enumerate(headers2, start=1):
            cell = ws2.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        curr_r2 = 4
        workflow_summary = plan_data.get("summary", [])
        for item in workflow_summary:
            row_vals = [
                item["serial_no"],
                item["section_name"],
                item["required_quantity"],
                item["total_target"],
                item["complete"],
                item["final_on_hand"],
                item["balance"],
                item["status"]
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws2.cell(row=curr_r2, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in [3, 4, 5, 6, 7]:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 8:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    stat_key = str(val).upper()
                    if stat_key in status_fills:
                        cell.fill = status_fills[stat_key]
            curr_r2 += 1

        autofit(ws2)

        # -------------------------------------------------------------
        # SHEET 3: Factory Calendar
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Factory Calendar")
        ws3["A1"] = "FACTORY WORKING CALENDAR & EXCEPTION SCHEDULE"
        ws3["A1"].font = title_font

        cal = CalendarService.get_active_calendar(db)
        schedule_rows = [
            ("Monday", "WORKING" if cal.work_monday else "OFF"),
            ("Tuesday", "WORKING" if cal.work_tuesday else "OFF"),
            ("Wednesday", "WORKING" if cal.work_wednesday else "OFF"),
            ("Thursday", "WORKING" if cal.work_thursday else "OFF"),
            ("Friday", "OFF (Default Bangladesh Weekend)" if not cal.work_friday else "WORKING"),
            ("Saturday", "WORKING" if cal.work_saturday else "OFF"),
            ("Sunday", "WORKING" if cal.work_sunday else "OFF"),
            ("Cutoff Time", f"{cal.cutoff_hour:02d}:{cal.cutoff_minute:02d}"),
            ("Timezone", cal.timezone),
        ]
        for idx, (k, v) in enumerate(schedule_rows, start=3):
            c_k = ws3.cell(row=idx, column=1, value=k)
            c_v = ws3.cell(row=idx, column=2, value=v)
            c_k.font = Font(bold=True)
            c_k.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
            c_k.border = thin_border
            c_v.border = thin_border

        # Calendar Exceptions Table
        ws3.cell(row=14, column=1, value="FACTORY HOLIDAYS & SPECIAL WORKING DAYS").font = Font(bold=True, color="1F4E79")
        ex_headers = ["Date", "Type", "Description"]
        for col_idx, h in enumerate(ex_headers, start=1):
            cell = ws3.cell(row=15, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        curr_r3 = 16
        for ex in cal.exceptions:
            e_vals = [str(ex.exception_date), str(ex.exception_type), ex.description]
            for col_idx, val in enumerate(e_vals, start=1):
                cell = ws3.cell(row=curr_r3, column=col_idx, value=val)
                cell.border = thin_border
            curr_r3 += 1

        autofit(ws3)

        # -------------------------------------------------------------
        # SHEET 4: Production Entries
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="Production Entries")
        ws4["A1"] = f"HISTORICAL PRODUCTION ENTRIES — {order.order_code}"
        ws4["A1"].font = title_font

        headers4 = ["Date", "Section", "Operation", "Accepted Qty", "Rejected Qty", "Recorded At", "Remarks"]
        ws4.row_dimensions[3].height = 22
        for col_idx, h in enumerate(headers4, start=1):
            cell = ws4.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        all_entries = db.query(DailyProductionEntry).filter(
            DailyProductionEntry.order_id == order.id,
            DailyProductionEntry.is_reversed == False
        ).order_by(DailyProductionEntry.entry_date.asc()).all()

        curr_r4 = 4
        for e in all_entries:
            sec_name = e.route.section.name if (e.route and e.route.section) else "N/A"
            op_name = e.operation.operation_name if e.operation else "N/A"
            entry_vals = [
                str(e.entry_date),
                sec_name,
                op_name,
                e.accepted_quantity,
                e.rejected_quantity,
                str(e.created_at.strftime("%Y-%m-%d %H:%M") if e.created_at else ""),
                e.remarks or ""
            ]
            for col_idx, val in enumerate(entry_vals, start=1):
                cell = ws4.cell(row=curr_r4, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in [4, 5]:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            curr_r4 += 1

        autofit(ws4)

        # -------------------------------------------------------------
        # SHEET 5: Order Details
        # -------------------------------------------------------------
        ws5 = wb.create_sheet(title="Order Details")
        ws5["A1"] = f"FOOTWEAR ORDER MASTER DETAILS — {order.order_code}"
        ws5["A1"].font = title_font

        comp_summary = plan_data.get("completion_summary", {})
        details_rows = [
            ("Order Code", order.order_code),
            ("Buyer", order.buyer.name if order.buyer else "N/A"),
            ("Customer Reference", order.customer_ref or "N/A"),
            ("Purchase Order (PO)", order.po_number or "N/A"),
            ("Product Style / Model", order.style_model or "N/A"),
            ("Description", order.description or "N/A"),
            ("Total Order Quantity (Pairs)", order.total_quantity),
            ("Planning Start Date", str(order.production_start_date)),
            ("Shipment / Delivery Date", str(order.current_delivery_date)),
            ("Remaining Working Days", comp_summary.get("remaining_working_days", 0)),
            ("Final Order Balance (Pairs)", comp_summary.get("final_balance", 0.0)),
            ("Last Section Cumulative Completed", comp_summary.get("completed_quantity", 0.0)),
            ("Current Bottleneck Process", comp_summary.get("current_bottleneck", "None")),
            ("WIP Waiting at Bottleneck", comp_summary.get("wip_waiting", 0.0)),
            ("Overall Production Status", comp_summary.get("overall_status", "N/A")),
        ]

        for idx, (k, v) in enumerate(details_rows, start=3):
            c_k = ws5.cell(row=idx, column=1, value=k)
            c_v = ws5.cell(row=idx, column=2, value=v)
            c_k.font = Font(bold=True)
            c_k.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
            c_k.border = thin_border
            c_v.border = thin_border
            if isinstance(v, (int, float)) and idx not in [10]:
                c_v.number_format = "#,##0"

        autofit(ws5)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    def generate_wip_flow_csv(cls, db: Session, order: Order) -> str:
        """
        Generates CSV format of the WIP Flow Planning Sheet.
        """
        from app.services.wip_planning_engine import WIPPlanningEngine
        plan_data = WIPPlanningEngine.calculate_order_wip_flow(db, order)

        output = io.StringIO()
        writer = csv.writer(output)

        # Header metadata
        writer.writerow(["FOOTWEAR MANUFACTURING WIP FLOW PLANNING SHEET"])
        writer.writerow(["Order Code", order.order_code, "Buyer", plan_data["buyer_name"]])
        writer.writerow(["Order Quantity", order.total_quantity, "Shipment Date", str(order.current_delivery_date)])
        writer.writerow(["Bottleneck", plan_data["current_bottleneck"], "WIP Waiting", plan_data["wip_waiting"]])
        writer.writerow([])

        # Table header
        headers = [
            "Date", "Serial", "Section", "T/G", "Opening On Hand", "Incoming WIP",
            "Actual", "Complete", "Closing On Hand", "Balance", "Remaining Working Days", "Status"
        ]
        writer.writerow(headers)

        daily_rows = plan_data.get("daily_plan_rows", [])
        for r in daily_rows:
            writer.writerow([
                str(r["date"]),
                r["serial_no"],
                r["section_name"],
                r["target"],
                r["opening_on_hand"],
                r["incoming_wip"],
                r["actual"],
                r["complete"],
                r["closing_on_hand"],
                r["balance"],
                r["remaining_working_days"],
                r["status"]
            ])

        writer.writerow([])
        writer.writerow(["WORKFLOW SUMMARY"])
        writer.writerow(["S/N", "Section", "Required Qty", "Total T/G", "Complete", "Final On Hand", "Balance", "Status"])
        for item in plan_data.get("summary", []):
            writer.writerow([
                item["serial_no"],
                item["section_name"],
                item["required_quantity"],
                item["total_target"],
                item["complete"],
                item["final_on_hand"],
                item["balance"],
                item["status"]
            ])

        return output.getvalue()


